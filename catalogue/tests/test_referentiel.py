"""Import du référentiel Excel : sessions, spécialités, matières, coefficients, chapitres."""

from datetime import date, datetime, timedelta

import pytest
from django.core.management import CommandError, call_command
from django.utils import timezone
from openpyxl import load_workbook

from catalogue.models import Chapitre, Examen, Matiere, Session, Specialite, SpecialiteMatiere
from catalogue.referentiel import ErreurClasseur, importer, modele
from sujets.models import Sujet

pytestmark = pytest.mark.django_db


@pytest.fixture
def remplir(tmp_path):
    """Écrit le modèle vide, y ajoute les lignes données par feuille et renvoie le chemin."""
    def _remplir(**feuilles):
        chemin = tmp_path / "referentiel.xlsx"
        modele(chemin)
        classeur = load_workbook(chemin)
        for nom, lignes in feuilles.items():
            for ligne in lignes:
                classeur["Spécialités" if nom == "specialites" else nom].append(ligne)
        classeur.save(chemin)
        return chemin
    return _remplir


def complet(remplir):
    return remplir(
        Sessions=[["BTS", 2025, datetime(2025, 4, 14), datetime(2025, 6, 27)], ["HND", 2025, "05/05/2025", "30/05/2025"]],
        specialites=[["BTS", "cge", "Comptabilité et gestion des entreprises", "Accounting", "Commerciale", 1],
                     ["HND", "SWE", "Software Engineering", "Génie logiciel", "industrielle", ""]],
        Matières=[["BTS", "Comptabilité générale", "Financial Accounting"], ["HND", "Programming", "Programmation"]],
        Coefficients=[["BTS", "CGE", "comptabilité  générale", 4], ["HND", "SWE", "Programming", 5]],
        Chapitres=[["BTS", "Comptabilité générale", 1, "Les amortissements", "Depreciation"]],
    )


def test_le_modele_a_ses_feuilles_et_ses_listes(tmp_path):
    modele(tmp_path / "m.xlsx")
    classeur = load_workbook(tmp_path / "m.xlsx")
    assert classeur.sheetnames == ["Lisez-moi", "Sessions", "Spécialités", "Matières", "Coefficients", "Chapitres"]
    assert [c.value for c in classeur["Sessions"][1]] == ["Examen", "Année", "Date de début", "Date de fin officielle"]
    assert any('"BTS,HND"' in v.formula1 for v in classeur["Matières"].data_validations.dataValidation)
    # Le modèle ne contient aucune donnée : un modèle importé tel quel ne crée rien.
    assert importer(tmp_path / "m.xlsx", appliquer=True).comptes == {}


def test_import_complet(remplir):
    rapport = importer(complet(remplir), appliquer=True)

    assert rapport.erreurs == []
    bts = Examen.objects.get(code="BTS")
    assert bts.langue == "fr" and Examen.objects.get(code="HND").langue == "en"
    session = Session.objects.get(examen=bts, annee=2025)
    assert session.date_fin_officielle == date(2025, 6, 27)
    assert Session.objects.get(examen__code="HND").date_fin_officielle == date(2025, 5, 30)  # date écrite en texte
    cge = Specialite.objects.get(code="CGE")
    assert cge.nom_fr == "Comptabilité et gestion des entreprises" and cge.nom_en == "Accounting" and cge.ordre == 1
    swe = Specialite.objects.get(code="SWE")
    assert swe.nom_en == "Software Engineering" and swe.nom_fr == "Génie logiciel" and swe.ordre == 100
    assert Matiere.objects.get(nom_en="Programming").nom_fr == "Programmation"
    # « comptabilité  générale » (casse, espaces) retrouve bien la matière.
    assert SpecialiteMatiere.objects.get(specialite=cge).coefficient == 4
    assert Chapitre.objects.get().titre_en == "Depreciation"
    assert rapport.comptes["Matières"] == {"créées": 2, "modifiées": 0, "inchangées": 0}


def test_essai_a_blanc_par_defaut(remplir):
    rapport = importer(complet(remplir))
    assert rapport.erreurs == [] and rapport.comptes["Sessions"]["créées"] == 2
    assert not Session.objects.exists() and not Matiere.objects.exists()


def test_relancer_le_meme_fichier_ne_double_rien(remplir):
    chemin = complet(remplir)
    importer(chemin, appliquer=True)
    rapport = importer(chemin, appliquer=True)
    assert Matiere.objects.count() == 2 and Session.objects.count() == 2 and SpecialiteMatiere.objects.count() == 2
    assert all(comptes["créées"] == comptes["modifiées"] == 0 for comptes in rapport.comptes.values())


def test_une_mise_a_jour_garde_les_traductions_non_remplies(remplir, compta_generale):
    compta_generale.nom_en = "Financial Accounting"
    compta_generale.save()
    importer(remplir(Matières=[["BTS", "Comptabilité générale", ""]]), appliquer=True)
    compta_generale.refresh_from_db()
    assert compta_generale.nom_en == "Financial Accounting"


def test_une_seule_erreur_et_rien_n_est_enregistre(remplir):
    rapport = importer(remplir(
        Sessions=[["BTS", 2025, "14/04/2025", "27/06/2025"], ["BTS", 2024, "30/06/2024", "01/06/2024"]],
        Matières=[["BTS", "Comptabilité générale", ""]],
    ), appliquer=True)
    assert rapport.erreurs == ["Sessions, ligne 3 : la date de fin officielle est avant la date de début."]
    assert not Session.objects.exists() and not Matiere.objects.exists()


@pytest.mark.parametrize("feuille, ligne, attendu", [
    ("Sessions", ["BAC", 2025, "14/04/2025", "27/06/2025"], "examen « BAC » inconnu"),
    ("Sessions", ["BTS", 2025, "14/04/2025", "fin juin"], "dates obligatoires"),
    ("Sessions", ["BTS", "deux mille", "14/04/2025", "27/06/2025"], "année invalide"),
    ("specialites", ["BTS", "CGE", "Comptabilité", "", "Tertiaire", 1], "filière"),
    ("Coefficients", ["BTS", "XYZ", "Inconnue", 30], "spécialité XYZ absente"),
    ("Chapitres", ["BTS", "Inconnue", 1, "Titre", ""], "matière « Inconnue » absente"),
])
def test_les_erreurs_donnent_la_feuille_et_la_ligne(remplir, feuille, ligne, attendu):
    rapport = importer(remplir(**{feuille: [ligne]}), appliquer=True)
    assert any(attendu in erreur for erreur in rapport.erreurs), rapport.erreurs
    assert all(", ligne 2 : " in erreur for erreur in rapport.erreurs)


def test_doublon_dans_le_fichier(remplir):
    rapport = importer(remplir(Matières=[["BTS", "Anglais", ""], ["BTS", "anglais", ""]]), appliquer=True)
    assert rapport.erreurs == ["Matières, ligne 3 : matière « anglais » écrite deux fois pour le BTS."]


def test_coefficient_d_une_matiere_d_un_autre_examen(remplir):
    """La matière est cherchée dans l'examen de la ligne : pas de coefficient BTS sur une matière HND."""
    rapport = importer(remplir(
        specialites=[["BTS", "IG", "Informatique de gestion", "", "Industrielle", 1]],
        Matières=[["HND", "Programming", ""]],
        Coefficients=[["BTS", "IG", "Programming", 4]],
    ), appliquer=True)
    assert rapport.erreurs == ["Coefficients, ligne 2 : matière « Programming » absente de la feuille Matières (BTS)."]


def test_repousser_la_fin_d_une_session_deja_publiee_est_refuse(remplir, creer_sujet, session_terminee):
    """Règle absolue : des sujets publiés ne peuvent pas se retrouver « publiés avant la fin officielle »."""
    creer_sujet(session_terminee, statut=Sujet.Statut.PUBLIE)
    plus_tard = timezone.localdate() + timedelta(days=10)
    rapport = importer(remplir(Sessions=[["BTS", 2025, "14/04/2025", plus_tard.strftime("%d/%m/%Y")]]), appliquer=True)

    assert len(rapport.erreurs) == 1 and "masquez d'abord ces sujets" in rapport.erreurs[0]
    session_terminee.refresh_from_db()
    assert session_terminee.date_fin_officielle == date(2025, 6, 27)


def test_corriger_une_fin_deja_passee_previent(remplir, creer_sujet, session_terminee):
    creer_sujet(session_terminee, statut=Sujet.Statut.PUBLIE)
    rapport = importer(remplir(Sessions=[["BTS", 2025, "14/04/2025", "30/06/2025"]]), appliquer=True)
    assert rapport.erreurs == []
    assert rapport.avertissements == [
        "BTS 2025 : date de fin officielle 27/06/2025 → 30/06/2025 (1 sujet(s) publié(s) dans cette session). "
        "Vérifiez la date sur le communiqué officiel."
    ]


def test_fichier_qui_n_est_pas_le_modele(tmp_path):
    autre = tmp_path / "autre.xlsx"
    from openpyxl import Workbook
    Workbook().save(autre)
    with pytest.raises(ErreurClasseur, match="Feuille « Sessions » introuvable"):
        importer(autre)
    (tmp_path / "abime.xlsx").write_bytes(b"pas un classeur")
    with pytest.raises(ErreurClasseur, match="Impossible de lire"):
        importer(tmp_path / "abime.xlsx")


def test_la_commande(remplir, capsys):
    chemin = complet(remplir)
    call_command("importer_referentiel", str(chemin))
    assert "Essai à blanc" in capsys.readouterr().out and not Session.objects.exists()
    call_command("importer_referentiel", str(chemin), "--appliquer")
    assert "Référentiel enregistré" in capsys.readouterr().out and Session.objects.count() == 2


def test_la_commande_refuse_un_fichier_faux(remplir):
    with pytest.raises(CommandError, match="1 erreur"):
        call_command("importer_referentiel", str(remplir(Matières=[["BAC", "Anglais", ""]])), "--appliquer")


# --- Équivalences BTS ↔ HND (CDC 3.10), colonne facultative ---------------------------------------------------

def test_le_modele_propose_la_colonne_equivalent(tmp_path):
    modele(tmp_path / "m.xlsx")
    assert [c.value for c in load_workbook(tmp_path / "m.xlsx")["Spécialités"][1]][-1] == "Équivalent (code dans l'autre examen)"


def test_l_equivalent_se_saisit_dans_le_classeur(remplir):
    chemin = remplir(specialites=[
        ["BTS", "CGE", "Comptabilité et gestion des entreprises", "", "Commerciale", 1, "acc"],
        ["HND", "ACC", "Accountancy", "", "Commerciale", 1, ""],  # une seule des deux lignes suffit
    ])
    rapport = importer(chemin, appliquer=True)
    assert rapport.erreurs == [] and rapport.comptes["Équivalences BTS ↔ HND"]["créées"] == 1
    cge, acc = Specialite.objects.get(code="CGE"), Specialite.objects.get(code="ACC")
    assert cge.equivalent == acc and acc.equivalent == cge
    assert importer(chemin, appliquer=True).comptes["Équivalences BTS ↔ HND"]["inchangées"] == 1


def test_un_classeur_sans_la_colonne_reste_valable(tmp_path):
    """Classeur commencé avant l'ajout de la colonne : rien n'est exigé."""
    from openpyxl import Workbook

    from catalogue.referentiel import FEUILLES

    classeur = Workbook()
    classeur.remove(classeur.active)
    for nom, colonnes in FEUILLES.items():
        classeur.create_sheet(nom).append(colonnes)
    classeur["Spécialités"].append(["BTS", "CGE", "Comptabilité et gestion des entreprises", "", "Commerciale", 1])
    classeur.save(tmp_path / "ancien.xlsx")
    assert importer(tmp_path / "ancien.xlsx", appliquer=True).erreurs == []


@pytest.mark.parametrize("lignes, attendu", [
    ([["BTS", "CGE", "Compta", "", "Commerciale", 1, "XYZ"]], "équivalent XYZ introuvable"),
    ([["BTS", "CGE", "Compta", "", "Commerciale", 1, "ACC"], ["BTS", "BF", "Banque", "", "Commerciale", 2, "ACC"],
      ["HND", "ACC", "Accountancy", "", "Commerciale", 1, ""]], "une spécialité n'a qu'une équivalente"),
    ([["BTS", "CGE", "Compta", "", "Commerciale", 1, "BF"], ["BTS", "BF", "Banque", "", "Commerciale", 2, ""]],
     "équivalent BF introuvable dans l'autre examen"),
])
def test_equivalent_refuse(remplir, lignes, attendu):
    rapport = importer(remplir(specialites=lignes), appliquer=True)
    assert any(attendu in erreur for erreur in rapport.erreurs), rapport.erreurs
    assert not Specialite.objects.exists()  # une erreur : rien n'est enregistré
