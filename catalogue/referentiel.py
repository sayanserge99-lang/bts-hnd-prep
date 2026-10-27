"""Référentiel des examens dans un classeur Excel : sessions, spécialités, matières, coefficients, chapitres.

Serge Alain remplit le modèle (`modele_referentiel`) dans Excel ou LibreOffice, puis `importer_referentiel`
le vérifie et l'enregistre. L'import ne supprime jamais rien : il crée ce qui manque et met à jour le reste,
on peut donc le relancer avec le même fichier. Une seule erreur et rien n'est enregistré.

Les noms suivent la langue de l'examen : pour le BTS (français), « Nom » est le nom français et « Nom dans
l'autre langue » l'anglais ; pour le HND (anglais), c'est l'inverse.
"""

import re
from dataclasses import dataclass, field
from datetime import date, datetime

from django.core.exceptions import ValidationError
from django.db import transaction
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from sujets.models import Sujet

from .models import Chapitre, Examen, Matiere, Session, Specialite, SpecialiteMatiere

NOMS_EXAMENS = {
    "BTS": ("Brevet de Technicien Supérieur", "fr"),
    "HND": ("Higher National Diploma", "en"),
}

# Feuille → colonnes, dans l'ordre du classeur.
FEUILLES = {
    "Sessions": ["Examen", "Année", "Date de début", "Date de fin officielle"],
    "Spécialités": ["Examen", "Code", "Nom", "Nom dans l'autre langue", "Filière", "Ordre d'affichage"],
    "Matières": ["Examen", "Nom", "Nom dans l'autre langue"],
    "Coefficients": ["Examen", "Code de la spécialité", "Matière", "Coefficient"],
    "Chapitres": ["Examen", "Matière", "Ordre", "Titre", "Titre dans l'autre langue"],
}

FILIERES = {"commerciale": "commerciale", "industrielle": "industrielle", "sante": "sante", "santé": "sante"}


class ErreurClasseur(Exception):
    """Le fichier n'est pas un classeur du modèle (feuille ou colonne manquante)."""


@dataclass
class Rapport:
    erreurs: list = field(default_factory=list)
    avertissements: list = field(default_factory=list)
    comptes: dict = field(default_factory=dict)

    def compter(self, feuille, quoi):
        self.comptes.setdefault(feuille, {"créées": 0, "modifiées": 0, "inchangées": 0})[quoi] += 1


# --- Lecture des cellules -------------------------------------------------------------------------------

def texte(valeur):
    if valeur is None:
        return ""
    if isinstance(valeur, float) and valeur.is_integer():
        valeur = int(valeur)
    return re.sub(r"\s+", " ", str(valeur)).strip()


def entier(valeur, mini, maxi):
    try:
        nombre = float(texte(valeur).replace(",", "."))
    except ValueError:
        return None
    if not nombre.is_integer() or not mini <= nombre <= maxi:
        return None
    return int(nombre)


def jour(valeur):
    """Une date Excel, ou un texte « 27/06/2025 » (ou « 2025-06-27 »)."""
    if isinstance(valeur, datetime):
        return valeur.date()
    if isinstance(valeur, date):
        return valeur
    brut = texte(valeur)
    for format_ in ("%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(brut, format_).date()
        except ValueError:
            pass
    return None


def lignes(classeur, nom):
    """(numéro de ligne Excel, {colonne: valeur}) pour chaque ligne remplie ; les lignes vides sont sautées."""
    if nom not in classeur.sheetnames:
        raise ErreurClasseur(f"Feuille « {nom} » introuvable : partez du modèle (python manage.py modele_referentiel).")
    feuille = classeur[nom]
    entete = [texte(cellule.value) for cellule in feuille[1]]
    manquantes = [colonne for colonne in FEUILLES[nom] if colonne not in entete]
    if manquantes:
        raise ErreurClasseur(f"Feuille « {nom} » : colonne(s) manquante(s) : {', '.join(manquantes)}.")
    for numero, valeurs in enumerate(feuille.iter_rows(min_row=2, values_only=True), start=2):
        ligne = {colonne: valeurs[entete.index(colonne)] if entete.index(colonne) < len(valeurs) else None
                 for colonne in FEUILLES[nom]}
        if any(texte(valeur) for valeur in ligne.values()):
            yield numero, ligne


def traduction(champ, langue, valeur):
    """La traduction n'est écrite que si la case est remplie : une case vide garde celle déjà saisie."""
    return {f"{champ}_{langue}": valeur} if valeur else {}


# --- Import ----------------------------------------------------------------------------------------------

class Import:
    def __init__(self):
        self.rapport = Rapport()
        self.examens = {}

    def erreur(self, feuille, numero, message):
        self.rapport.erreurs.append(f"{feuille}, ligne {numero} : {message}")

    def examen(self, feuille, numero, valeur):
        code = texte(valeur).upper()
        if code not in NOMS_EXAMENS:
            self.erreur(feuille, numero, f"examen « {texte(valeur)} » inconnu (BTS ou HND).")
            return None
        if code not in self.examens:
            nom, langue = NOMS_EXAMENS[code]
            self.examens[code], _ = Examen.objects.get_or_create(code=code, defaults={"nom": nom, "langue": langue})
        return self.examens[code]

    @staticmethod
    def langues(examen):
        """(langue principale, autre langue) de l'examen."""
        return ("fr", "en") if examen.langue == "fr" else ("en", "fr")

    def enregistrer(self, feuille, objet, valeurs, cree):
        """Applique `valeurs` à un objet existant et dit s'il a changé."""
        if cree:
            self.rapport.compter(feuille, "créées")
            return
        changes = {nom: valeur for nom, valeur in valeurs.items() if getattr(objet, nom) != valeur}
        if not changes:
            self.rapport.compter(feuille, "inchangées")
            return
        for nom, valeur in changes.items():
            setattr(objet, nom, valeur)
        objet.save()
        self.rapport.compter(feuille, "modifiées")

    def sessions(self, classeur):
        vues = set()
        for numero, ligne in lignes(classeur, "Sessions"):
            examen = self.examen("Sessions", numero, ligne["Examen"])
            annee = entier(ligne["Année"], 2000, 2100)
            debut, fin = jour(ligne["Date de début"]), jour(ligne["Date de fin officielle"])
            if annee is None:
                self.erreur("Sessions", numero, "année invalide (exemple : 2025).")
            if debut is None or fin is None:
                self.erreur("Sessions", numero, "dates obligatoires, au format JJ/MM/AAAA.")
            elif fin < debut:
                self.erreur("Sessions", numero, "la date de fin officielle est avant la date de début.")
            if not examen or annee is None or debut is None or fin is None or fin < debut:
                continue
            if (examen.code, annee) in vues:
                self.erreur("Sessions", numero, f"session {examen.code} {annee} écrite deux fois.")
                continue
            vues.add((examen.code, annee))
            session = Session.objects.filter(examen=examen, annee=annee).first()
            if session is None:
                Session.objects.create(examen=examen, annee=annee, date_debut=debut, date_fin_officielle=fin)
                self.rapport.compter("Sessions", "créées")
                continue
            ancienne_fin = session.date_fin_officielle
            try:
                # Session.save refuse une fin à venir si des sujets sont déjà publiés (règle absolue).
                self.enregistrer("Sessions", session, {"date_debut": debut, "date_fin_officielle": fin}, False)
            except ValidationError as erreur:
                self.erreur("Sessions", numero, f"{session} : {' '.join(erreur.messages)}")
                continue
            if session.date_fin_officielle != ancienne_fin:
                publies = Sujet.objects.filter(session=session, statut=Sujet.Statut.PUBLIE).count()
                self.rapport.avertissements.append(
                    f"{session} : date de fin officielle {ancienne_fin:%d/%m/%Y} → {fin:%d/%m/%Y}"
                    + (f" ({publies} sujet(s) publié(s) dans cette session)" if publies else "")
                    + ". Vérifiez la date sur le communiqué officiel."
                )

    def specialites(self, classeur):
        vues = set()
        for numero, ligne in lignes(classeur, "Spécialités"):
            examen = self.examen("Spécialités", numero, ligne["Examen"])
            code = texte(ligne["Code"]).upper()
            nom, autre = texte(ligne["Nom"]), texte(ligne["Nom dans l'autre langue"])
            filiere = FILIERES.get(texte(ligne["Filière"]).lower())
            ordre = entier(ligne["Ordre d'affichage"], 0, 999) if texte(ligne["Ordre d'affichage"]) else 100
            if not code or len(code) > 20:
                self.erreur("Spécialités", numero, "code obligatoire, 20 caractères au plus (exemple : CGE).")
            if not nom or len(nom) > 200 or len(autre) > 200:
                self.erreur("Spécialités", numero, "nom obligatoire, 200 caractères au plus.")
            if filiere is None:
                self.erreur("Spécialités", numero, "filière : Commerciale, Industrielle ou Santé.")
            if ordre is None:
                self.erreur("Spécialités", numero, "ordre d'affichage : un nombre entier (plus petit = en premier).")
            if not examen or not code or len(code) > 20 or not nom or len(nom) > 200 or len(autre) > 200                     or filiere is None or ordre is None:
                continue
            if (examen.code, code) in vues:
                self.erreur("Spécialités", numero, f"code {code} écrit deux fois pour le {examen.code}.")
                continue
            vues.add((examen.code, code))
            principale, seconde = self.langues(examen)
            valeurs = {f"nom_{principale}": nom, **traduction("nom", seconde, autre), "filiere": filiere, "ordre": ordre}
            objet, cree = Specialite.objects.get_or_create(examen=examen, code=code, defaults=valeurs)
            self.enregistrer("Spécialités", objet, valeurs, cree)

    def trouver_matiere(self, examen, nom):
        principale, _ = self.langues(examen)
        return Matiere.objects.filter(examen=examen, **{f"nom_{principale}__iexact": nom}).first()

    def matieres(self, classeur):
        vues = set()
        for numero, ligne in lignes(classeur, "Matières"):
            examen = self.examen("Matières", numero, ligne["Examen"])
            nom, autre = texte(ligne["Nom"]), texte(ligne["Nom dans l'autre langue"])
            if not nom or len(nom) > 200 or len(autre) > 200:
                self.erreur("Matières", numero, "nom obligatoire, 200 caractères au plus.")
            if not examen or not nom or len(nom) > 200 or len(autre) > 200:
                continue
            if (examen.code, nom.lower()) in vues:
                self.erreur("Matières", numero, f"matière « {nom} » écrite deux fois pour le {examen.code}.")
                continue
            vues.add((examen.code, nom.lower()))
            principale, seconde = self.langues(examen)
            valeurs = {f"nom_{principale}": nom, **traduction("nom", seconde, autre)}
            objet = self.trouver_matiere(examen, nom)
            if objet is None:
                Matiere.objects.create(examen=examen, **valeurs)
                self.rapport.compter("Matières", "créées")
            else:
                self.enregistrer("Matières", objet, valeurs, False)

    def coefficients(self, classeur):
        for numero, ligne in lignes(classeur, "Coefficients"):
            examen = self.examen("Coefficients", numero, ligne["Examen"])
            if not examen:
                continue
            code, nom = texte(ligne["Code de la spécialité"]).upper(), texte(ligne["Matière"])
            coefficient = entier(ligne["Coefficient"], 1, 20)
            specialite = Specialite.objects.filter(examen=examen, code=code).first()
            matiere = self.trouver_matiere(examen, nom)
            if specialite is None:
                self.erreur("Coefficients", numero, f"spécialité {code or '(vide)'} absente de la feuille Spécialités ({examen.code}).")
            if matiere is None:
                self.erreur("Coefficients", numero, f"matière « {nom} » absente de la feuille Matières ({examen.code}).")
            if coefficient is None:
                self.erreur("Coefficients", numero, "coefficient : un nombre entier de 1 à 20.")
            if specialite is None or matiere is None or coefficient is None:
                continue
            objet, cree = SpecialiteMatiere.objects.get_or_create(
                specialite=specialite, matiere=matiere, defaults={"coefficient": coefficient}
            )
            self.enregistrer("Coefficients", objet, {"coefficient": coefficient}, cree)

    def chapitres(self, classeur):
        for numero, ligne in lignes(classeur, "Chapitres"):
            examen = self.examen("Chapitres", numero, ligne["Examen"])
            if not examen:
                continue
            nom, titre, autre = texte(ligne["Matière"]), texte(ligne["Titre"]), texte(ligne["Titre dans l'autre langue"])
            ordre = entier(ligne["Ordre"], 0, 999) if texte(ligne["Ordre"]) else 0
            matiere = self.trouver_matiere(examen, nom)
            if matiere is None:
                self.erreur("Chapitres", numero, f"matière « {nom} » absente de la feuille Matières ({examen.code}).")
            if not titre or len(titre) > 200 or len(autre) > 200:
                self.erreur("Chapitres", numero, "titre obligatoire, 200 caractères au plus.")
            if ordre is None:
                self.erreur("Chapitres", numero, "ordre : un nombre entier.")
            if matiere is None or not titre or len(titre) > 200 or len(autre) > 200 or ordre is None:
                continue
            principale, seconde = self.langues(examen)
            valeurs = {f"titre_{principale}": titre, **traduction("titre", seconde, autre), "ordre": ordre}
            objet = Chapitre.objects.filter(matiere=matiere, **{f"titre_{principale}__iexact": titre}).first()
            if objet is None:
                Chapitre.objects.create(matiere=matiere, **valeurs)
                self.rapport.compter("Chapitres", "créées")
            else:
                self.enregistrer("Chapitres", objet, valeurs, False)


def importer(fichier, appliquer=False):
    """Lit le classeur et renvoie le rapport. Sans `appliquer`, ou s'il y a une erreur, rien n'est gardé."""
    try:
        classeur = load_workbook(fichier, data_only=True, read_only=True)
    except Exception as exc:  # fichier abîmé, pas un .xlsx…
        raise ErreurClasseur(f"Impossible de lire le classeur ({exc.__class__.__name__}) : enregistrez-le au format .xlsx.")
    travail = Import()
    with transaction.atomic():
        # Dans l'ordre : les coefficients et les chapitres renvoient aux spécialités et matières déjà lues.
        for etape in (travail.sessions, travail.specialites, travail.matieres, travail.coefficients, travail.chapitres):
            etape(classeur)
        if travail.rapport.erreurs or not appliquer:
            transaction.set_rollback(True)
    return travail.rapport


# --- Modèle vide -----------------------------------------------------------------------------------------

LISEZ_MOI = [
    "Référentiel BTS-HND Prep : à remplir puis à importer (python manage.py importer_referentiel fichier.xlsx).",
    "",
    "Une ligne par élément. Ne renommez ni les feuilles ni les titres de colonnes.",
    "Examen : BTS ou HND. Les noms sont dans la langue de l'examen (BTS : français ; HND : anglais) ;",
    "« dans l'autre langue » est facultatif mais conseillé (le site est bilingue).",
    "",
    "Sessions — la date de fin officielle décide de la publication : aucun sujet de la session n'est publié",
    "avant le lendemain de cette date. Recopiez-la du communiqué officiel. Format JJ/MM/AAAA.",
    "Spécialités — le code est court et unique par examen (CGE, IG, SWE…). Filière : Commerciale, Industrielle",
    "ou Santé. Ordre d'affichage : plus petit = affiché en premier (les spécialités du lancement d'abord).",
    "Matières — une matière appartient à un seul examen (Comptabilité générale au BTS, Financial Accounting au HND).",
    "Coefficients — une ligne par couple spécialité / matière ; la matière est écrite comme dans la feuille Matières.",
    "Chapitres — facultatif ; sert aux quiz par chapitre.",
    "",
    "L'import ne supprime jamais rien et peut être relancé avec le même fichier : il met à jour ce qui a changé.",
    "Il vérifie d'abord tout le fichier ; à la moindre erreur, il donne la feuille et la ligne et n'enregistre rien.",
]

EXEMPLES = {
    "Sessions": ["BTS", 2025, "14/04/2025", "27/06/2025"],
    "Spécialités": ["BTS", "CGE", "Comptabilité et gestion des entreprises", "Accounting and Business Management",
                    "Commerciale", 1],
    "Matières": ["BTS", "Comptabilité générale", "Financial Accounting"],
    "Coefficients": ["BTS", "CGE", "Comptabilité générale", 4],
    "Chapitres": ["BTS", "Comptabilité générale", 1, "Les amortissements", "Depreciation"],
}


def modele(chemin):
    """Écrit le classeur vide, avec ses listes déroulantes et une ligne d'exemple sur la page Lisez-moi."""
    classeur = Workbook()
    lisez_moi = classeur.active
    lisez_moi.title = "Lisez-moi"
    for numero, phrase in enumerate(LISEZ_MOI, start=1):
        lisez_moi.cell(numero, 1, phrase)
    lisez_moi["A1"].font = Font(bold=True, size=13)
    lisez_moi.column_dimensions["A"].width = 115
    ligne = len(LISEZ_MOI) + 2
    lisez_moi.cell(ligne, 1, "Exemples (à recopier dans les feuilles, pas à laisser ici) :").font = Font(bold=True)
    for nom, exemple in EXEMPLES.items():
        ligne += 1
        lisez_moi.cell(ligne, 1, f"{nom} : " + " | ".join(str(valeur) for valeur in exemple))

    entete = PatternFill("solid", fgColor="DDEBDD")
    for nom, colonnes in FEUILLES.items():
        feuille = classeur.create_sheet(nom)
        feuille.append(colonnes)
        feuille.freeze_panes = "A2"
        for numero, colonne in enumerate(colonnes, start=1):
            cellule = feuille.cell(1, numero)
            cellule.font, cellule.fill = Font(bold=True), entete
            cellule.alignment = Alignment(wrap_text=True)
            lettre = cellule.column_letter
            feuille.column_dimensions[lettre].width = 45 if colonne.startswith(("Nom", "Matière", "Titre")) else 18
            plage = f"{lettre}2:{lettre}1000"
            if colonne == "Examen":
                liste = DataValidation(type="list", formula1='"BTS,HND"', allow_blank=True)
            elif colonne == "Filière":
                liste = DataValidation(type="list", formula1='"Commerciale,Industrielle,Santé"', allow_blank=True)
            elif colonne.startswith("Date"):
                feuille.column_dimensions[lettre].number_format = "DD/MM/YYYY"
                liste = DataValidation(type="date", operator="greaterThan", formula1="36526", allow_blank=True,
                                       error="Une date, au format JJ/MM/AAAA.", showErrorMessage=True)
            else:
                continue
            liste.add(plage)
            feuille.add_data_validation(liste)
    classeur.save(chemin)
